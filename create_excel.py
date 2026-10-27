import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

INPUT_FILE = "C:\\Users\\pedroc\\Documents\\swiss_houses\\vaud_apartments_final.csv"
OUTPUT_FILE = "C:\\Users\\pedroc\\Documents\\swiss_houses\\vaud_apartments.xlsx"


def main():
    # Read CSV
    rows = []
    headers = []
    with open(INPUT_FILE, 'r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        headers = reader.fieldnames
        for row in reader:
            rows.append(row)
    
    print(f"Read {len(rows)} rows")
    
    # Create workbook
    wb = openpyxl.Workbook()
    
    # Style definitions
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    price_font = Font(bold=True, color="006100")
    price_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Sheet 1: All listings
    ws1 = wb.active
    ws1.title = "All Listings"
    
    # Write headers
    for col, header in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    # Write data
    for row_idx, row in enumerate(rows, 2):
        for col_idx, header in enumerate(headers, 1):
            value = row.get(header, '')
            # Try to convert numeric values
            if header in ['Price (CHF)', 'Price/m2 (CHF)', 'Size (m2)', 'Year Built']:
                try:
                    value = int(value) if value else ''
                except:
                    try:
                        value = float(value) if value else ''
                    except:
                        pass
            
            cell = ws1.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            
            # Highlight price columns
            if header == 'Price (CHF)' and value:
                cell.font = price_font
                cell.fill = price_fill
                cell.number_format = '#,##0'
            elif header == 'Price/m2 (CHF)' and value:
                cell.font = price_font
                cell.fill = price_fill
                cell.number_format = '#,##0'
    
    # Auto-fit column widths
    for col in range(1, len(headers) + 1):
        max_length = 0
        column_letter = openpyxl.utils.get_column_letter(col)
        for row in range(1, len(rows) + 2):
            try:
                cell_value = str(ws1.cell(row=row, column=col).value or '')
                max_length = max(max_length, len(cell_value))
            except:
                pass
        adjusted_width = min(max_length + 2, 30)
        ws1.column_dimensions[column_letter].width = adjusted_width
    
    # Freeze top row
    ws1.freeze_panes = 'A2'
    
    # Sheet 2: Sorted by Price/m2
    ws2 = wb.create_sheet("By Price per m2")
    
    # Sort rows by Price/m2
    def get_price_per_m2(row):
        try:
            return int(row.get('Price/m2 (CHF)', '999999') or '999999')
        except:
            return 999999
    
    sorted_rows = sorted(rows, key=get_price_per_m2)
    
    # Filter out rows without Price/m2
    sorted_rows = [r for r in sorted_rows if r.get('Price/m2 (CHF)')]
    
    print(f"Sorted {len(sorted_rows)} rows by Price/m2")
    
    # Write headers (Rank column + data columns shifted by 1)
    rank_cell = ws2.cell(row=1, column=1, value="#")
    rank_cell.font = header_font
    rank_cell.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
    rank_cell.alignment = header_alignment
    rank_cell.border = thin_border
    
    for col, header in enumerate(headers, 2):
        cell = ws2.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
        cell.alignment = header_alignment
        cell.border = thin_border
    
    # Write sorted data with ranking
    for row_idx, row in enumerate(sorted_rows, 2):
        # Add rank column
        cell = ws2.cell(row=row_idx, column=1, value=row_idx - 1)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")
        
        for col_idx, header in enumerate(headers, 2):
            value = row.get(header, '')
            # Try to convert numeric values
            if header in ['Price (CHF)', 'Price/m2 (CHF)', 'Size (m2)', 'Year Built']:
                try:
                    value = int(value) if value else ''
                except:
                    try:
                        value = float(value) if value else ''
                    except:
                        pass
            
            cell = ws2.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            
            # Highlight price columns
            if header == 'Price (CHF)' and value:
                cell.font = price_font
                cell.fill = price_fill
                cell.number_format = '#,##0'
            elif header == 'Price/m2 (CHF)' and value:
                cell.font = price_font
                cell.fill = price_fill
                cell.number_format = '#,##0'
        
        # Color code by value (green = cheap, red = expensive)
        try:
            ppm = int(row.get('Price/m2 (CHF)', '0') or '0')
            if ppm > 0:
                if ppm < 10000:
                    fill_color = "C6EFCE"  # Light green
                elif ppm < 15000:
                    fill_color = "FFEB9C"  # Light yellow
                else:
                    fill_color = "FFC7CE"  # Light red
                
                for col_idx in range(1, len(headers) + 2):
                    ws2.cell(row=row_idx, column=col_idx).fill = PatternFill(
                        start_color=fill_color, end_color=fill_color, fill_type="solid"
                    )
        except:
            pass
    
    # Auto-fit column widths
    for col in range(1, len(headers) + 2):
        max_length = 0
        column_letter = openpyxl.utils.get_column_letter(col)
        for row in range(1, len(sorted_rows) + 2):
            try:
                cell_value = str(ws2.cell(row=row, column=col).value or '')
                max_length = max(max_length, len(cell_value))
            except:
                pass
        adjusted_width = min(max_length + 2, 30)
        ws2.column_dimensions[column_letter].width = adjusted_width
    
    # Freeze top row
    ws2.freeze_panes = 'A2'
    
    # Save
    wb.save(OUTPUT_FILE)
    print(f"\nSaved to {OUTPUT_FILE}")
    
    # Print summary of sorted list
    print("\nTop 10 Best Value (lowest Price/m2):")
    for i, row in enumerate(sorted_rows[:10], 1):
        ppm = row.get('Price/m2 (CHF)', '?')
        price = row.get('Price (CHF)', '?')
        city = row.get('City', '?')
        size = row.get('Size (m2)', '?')
        print(f"  {i}. {city} | CHF {price} | {size}m2 | CHF {ppm}/m2")


if __name__ == '__main__':
    main()
